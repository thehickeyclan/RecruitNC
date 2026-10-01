#!/usr/bin/env python3
"""
Import an out-of-state state-tournament workbook into Supabase
(state_tournament_divisions / state_tournament_bouts / state_tournament_placers).

Workbook layout (the collector template):
  events_and_divisions  one row per class: places awarded, class rank, dates
  matches               one row per placement bout (1st/3rd/5th/7th) — the only sheet trusted for names
  placers               read only for matching fields (school clean, city, grade, grad year, athlete id);
                        its name/school columns are formulas and are ignored — placers are derived from matches

The workbook is the full truth for every division it contains: those divisions' bouts and placers are
replaced on each run, so re-running with a more complete file (e.g. 5th/6th added) just fills in.

Prerequisite: run scripts/create-state-tournament-results.sql in the Supabase SQL Editor.

  python3 scripts/import-state-tournament-results.py "~/Downloads/Virginia 2026 Wrestling Results.xlsx" --dry-run
  python3 scripts/import-state-tournament-results.py "~/Downloads/Virginia 2026 Wrestling Results.xlsx"

After writing, it marks identity_confirmed on every placer that some NC bout on file already ties to
him - his school on the bout, or the bout listing his state. Re-run that step alone after importing
new national bouts, since a new bout can confirm a name:

  python3 scripts/import-state-tournament-results.py --confirm-only
"""

import importlib.util
import json
import os
import time
import re
import sys
import urllib.parse
import urllib.request
from collections import defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULT_TYPES = {"F", "TF", "MD", "DEC", "SV", "TB", "UTB", "INJ", "DQ", "FF"}
BOUT_PLACES = {"1st": (1, 2), "3rd": (3, 4), "5th": (5, 6), "7th": (7, 8)}
DIVISION_KEY = ("season", "state", "association", "gender", "classification")


def load_env():
    for name in (".env.local", ".env"):
        path = os.path.join(ROOT, name)
        if not os.path.exists(path):
            continue
        for line in open(path):
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, val = line.split("=", 1)
            val = val.strip().strip('"').strip("'")
            os.environ.setdefault(key.strip(), val)


def text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if not s or s.startswith("="):
        return None
    return s


def wrestler(raw):
    """A wrestler's name, or None where the sheet put a placeholder in the name column."""
    s = text(raw)
    return None if s and s.lower() in {"forfeit", "bye", "tbd", "unknown", "n/a"} else s


def normalize_class(raw):
    s = text(raw)
    m = re.fullmatch(r"Class\s+(\d)", s or "", re.I)
    return f"{m.group(1)}A" if m else s


def sheet_rows(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    header = [text(h) for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:]]


def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    errors = []

    divisions = {}
    for r in sheet_rows(wb, "events_and_divisions"):
        if not isinstance(r.get("Season"), (int, float)):
            continue  # blank spacer or the trailing NOTE row
        d = {
            "season": int(r["Season"]),
            "state": text(r["State"]),
            "association": text(r["Association"]),
            "gender": text(r["Gender"]),
            "classification": normalize_class(r["Classification"]),
            "class_rank": int(r["Class Rank (1 = largest)"]) if r.get("Class Rank (1 = largest)") else None,
            "classes_in_state": int(r["Classes in State"]) if r.get("Classes in State") else None,
            "brackets": int(r["Brackets"]) if r.get("Brackets") else None,
            "places_awarded": int(r["Places Awarded"]),
            "tournament_dates": text(r.get("Tournament Dates")),
            "source_url": text(r.get("Bracket Source URL")),
        }
        divisions[tuple(d[k] for k in DIVISION_KEY)] = d

    bouts = []
    for i, r in enumerate(sheet_rows(wb, "matches"), start=2):
        winner, loser = wrestler(r.get("Winner")), wrestler(r.get("Loser"))
        if not winner and not loser:
            continue  # bout row not collected yet
        b = {
            "season": int(r["Season"]),
            "state": text(r["State"]),
            "association": text(r["Association"]),
            "gender": text(r["Gender"]),
            "classification": normalize_class(r["Classification"]),
            "weight": text(r["Weight"]),
            "bout": text(r["Bout (1st/3rd/5th/7th)"]),
            "winner_name": winner,
            "winner_school": text(r.get("Winner School")),
            "loser_name": loser,
            "loser_school": text(r.get("Loser School")),
            "result_type": text(r.get("Result Type")),
            "score": text(r.get("Score")),
            "fall_time": text(r.get("Fall Time")),
            "source_url": text(r.get("Source URL")),
        }
        where = f"matches row {i} ({b['classification']} {b['weight']} {b['bout']})"
        div = divisions.get(tuple(b[k] for k in DIVISION_KEY))
        # A placement bout with no loser happens - the other wrestler forfeited or never made the
        # bout - and the winner still placed. Only a bout with no winner is unusable.
        if not winner:
            errors.append(f"{where}: winner missing")
        if b["bout"] not in BOUT_PLACES:
            errors.append(f"{where}: bout must be 1st/3rd/5th/7th")
        elif div and BOUT_PLACES[b["bout"]][1] > div["places_awarded"]:
            errors.append(f"{where}: bout deeper than places_awarded={div['places_awarded']}")
        if not div:
            errors.append(f"{where}: no events_and_divisions row for this class")
        if b["result_type"] and b["result_type"] not in RESULT_TYPES:
            errors.append(f"{where}: result type {b['result_type']!r} is not a fixed code")
        if b["result_type"] == "F" and not b["fall_time"]:
            errors.append(f"{where}: fall without a fall time")
        bouts.append(b)

    seen = defaultdict(int)
    for b in bouts:
        seen[tuple(b[k] for k in DIVISION_KEY) + (b["weight"], b["bout"])] += 1
    errors += [f"duplicate bout {k}" for k, n in seen.items() if n > 1]

    extras = {}
    if "placers" in wb.sheetnames:
        for r in sheet_rows(wb, "placers"):
            if not isinstance(r.get("Season"), (int, float)):
                continue
            key = (normalize_class(r["Classification"]), text(r["Weight"]), int(r["Place"]))
            grad = text(r.get("Grad Year"))
            extras[key] = {
                "school_clean": text(r.get("School (clean)")),
                "city": text(r.get("City")),
                "grade": text(r.get("Grade")),
                "grad_year": int(grad) if grad and grad.isdigit() else None,
                "source_athlete_id": text(r.get("Athlete ID")),
                "source_athlete_id_source": text(r.get("Athlete ID Source")),
            }

    placers = []
    for b in bouts:
        if b["bout"] not in BOUT_PLACES:
            continue
        win_place, lose_place = BOUT_PLACES[b["bout"]]
        for place, name, school in ((win_place, b["winner_name"], b["winner_school"]),
                                    (lose_place, b["loser_name"], b["loser_school"])):
            if not name:
                continue
            row = {k: b[k] for k in DIVISION_KEY + ("weight",)}
            row.update(place=place, wrestler_name=name, school_raw=school, source_url=b["source_url"])
            row.update(extras.get((b["classification"], b["weight"], place), {}))
            placers.append(row)

    return divisions, bouts, placers, errors


def coverage(divisions, bouts):
    lines = []
    for key, d in divisions.items():
        got = sum(1 for b in bouts if tuple(b[k] for k in DIVISION_KEY) == key)
        expected = (d["brackets"] or 0) * ((d["places_awarded"] + 1) // 2)
        lines.append(f"  {d['state']} {d['season']} {d['gender']} {d['classification']}: "
                     f"{got}/{expected} placement bouts, {got * 2}/{expected * 2} placers")
    return "\n".join(lines)


SCHOOL_STOP = {"high", "school", "hs", "senior", "the", "of", "academy", "christian"}


def school_words(value):
    """Mirror of schoolWords in lib/significant-wins.ts — the two must agree on what a match is."""
    return [w for w in re.split(r"\s+", re.sub(r"[^a-z0-9\s]", " ", str(value or "").lower())) if w and w not in SCHOOL_STOP]


def same_school(a, b):
    return bool(a) and bool(b) and (all(w in b for w in a) or all(w in a for w in b))


def name_key(value):
    return " ".join(re.sub(r"[^a-z ]", "", re.sub(r"[-'.]", " ", str(value or "").lower())).split())


def bout_confirms(placer, school):
    """The same evidence outOfStatePlacerFits accepts on its own: his state listed, or his school."""
    text = str(school or "").strip().upper()
    state = placer["state"].upper()
    if text == state or f"({state})" in text:
        return True
    return same_school(school_words(placer["school_raw"]), school_words(school))


# Placers confirmed by hand where the data cannot: a distinctive name on a bout that carries no
# state or school. Each line says why. The pass below always marks these confirmed.
MANUAL_CONFIRMED = {
    # Beat Jake Amiott at the 2026 NHSCA National Duals for Team Shutt Rioux, an all-star side with
    # no home state. Only Tyler Traves among 11,600 placers; wrestled 145 there, won VA 6A at 150.
    ("VA", "tyler traves"),
}


def confirm_identities(db):
    placers = db.get_all("state_tournament_placers", "select=*")
    by_name = defaultdict(list)
    for p in placers:
        by_name[name_key(p["wrestler_name"])].append(p)

    sightings = []  # (opponent name, school or club as the bout lists it)
    for b in db.get_all("other_tournament_bouts", "select=opponent_name,opponent_club&is_bye=eq.false"):
        sightings.append((b["opponent_name"], b["opponent_club"]))
    for r in db.get_all("matches", "select=matches", page=200):
        bouts = r["matches"] if isinstance(r["matches"], list) else json.loads(r["matches"] or "[]")
        sightings += [(m.get("opponent"), m.get("opponent_school")) for m in bouts]

    confirmed = {p["id"] for p in placers if (p["state"], name_key(p["wrestler_name"])) in MANUAL_CONFIRMED}
    for name, school in sightings:
        for p in by_name.get(name_key(name), []):
            if bout_confirms(p, school):
                confirmed.add(p["id"])

    # Only rows whose flag changes: rewriting all 11,600 every run exhausted local ports.
    for value in (True, False):
        changed = [p["id"] for p in placers if (p["id"] in confirmed) == value and bool(p.get("identity_confirmed")) != value]
        for chunk in range(0, len(changed), 200):
            db.request("PATCH", "state_tournament_placers", f"?id=in.({','.join(changed[chunk:chunk + 200])})",
                       {"identity_confirmed": value}, "return=minimal")
    names = {name_key(p["wrestler_name"]) for p in placers if p["id"] in confirmed}
    print(f"identity_confirmed: {len(names)} of {len(by_name)} placer names tied to their state by a bout on file")
    mark_nc_namesakes(db, placers)


NC_SCHOOL_STOP = {"high", "school", "hs", "senior", "the", "of", "academy", "christian"}


def mark_nc_namesakes(db, placers):
    """
    nc_namesake: someone with this placer's name - exactly, or by surname and a nickname - has
    wrestled for a North Carolina high school in our season records. The profile credits a
    distinctive out-of-state name on its own (outOfStatePlacerFits), and only when this is false:
    Mooresville's Gavin Walker is not Virginia's.
    """
    if "nc_namesake" not in placers[0]:
        print("nc_namesake: column missing - run the ALTER in scripts/create-state-tournament-results.sql")
        return
    words = lambda v: [w for w in name_key(v).split() if w not in NC_SCHOOL_STOP]
    nc_schools = [words(r["school"]) for r in db.get_all("wrestling_nchsaa_results", "select=school&year=gte.2018") if r.get("school")]
    nc_schools = [w for w in nc_schools if w]

    def at_nc_school(school):
        b = words(school)
        return bool(b) and any(all(x in b for x in a) or all(x in a for x in b) for a in nc_schools)

    seen = set()
    for r in db.get_all("matches", "select=matches", page=100):
        bouts = r["matches"] if isinstance(r["matches"], list) else json.loads(r["matches"] or "[]")
        for m in bouts:
            if m.get("opponent") and at_nc_school(m.get("opponent_school")):
                seen.add(name_key(m["opponent"]))
    for a in db.get_all("athletes", "select=name&is_nc_athlete=eq.true"):
        if a.get("name"):
            seen.add(name_key(a["name"]))

    spec = importlib.util.spec_from_file_location("fargo", os.path.join(ROOT, "scripts", "import-fargo-bouts.py"))
    fargo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fargo)
    groups = fargo.nickname_groups()
    by_surname = defaultdict(set)
    for n in seen:
        parts = n.split()
        if len(parts) >= 2:
            by_surname[parts[-1]].add(parts[0])

    def namesake(name):
        key = name_key(name)
        if key in seen:
            return True
        parts = key.split()
        return len(parts) >= 2 and any(fargo.first_names_alike(parts[0], f, groups) for f in by_surname.get(parts[-1], ()))

    for value in (True, False):
        changed = [p["id"] for p in placers if namesake(p["wrestler_name"]) == value and bool(p.get("nc_namesake")) != value]
        for chunk in range(0, len(changed), 200):
            db.request("PATCH", "state_tournament_placers", f"?id=in.({','.join(changed[chunk:chunk + 200])})",
                       {"nc_namesake": value}, "return=minimal")
    print(f"nc_namesake: {sum(1 for p in placers if namesake(p['wrestler_name']))} placers share a name with a North Carolinian")


class Supabase:
    def __init__(self):
        self.url = (os.environ.get("NEXT_PUBLIC_SUPABASE_URL") or os.environ.get("SUPABASE_URL") or "").rstrip("/")
        self.key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
        if not self.url or not self.key:
            sys.exit("Missing NEXT_PUBLIC_SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY in .env.local")

    def request(self, method, table, query="", body=None, prefer=None):
        req = urllib.request.Request(f"{self.url}/rest/v1/{table}{query}", method=method,
                                     data=json.dumps(body).encode() if body is not None else None)
        req.add_header("apikey", self.key)
        req.add_header("Authorization", f"Bearer {self.key}")
        req.add_header("Content-Type", "application/json")
        if prefer:
            req.add_header("Prefer", prefer)
        for attempt in range(5):
            try:
                with urllib.request.urlopen(req) as res:
                    return res.read()
            except urllib.error.HTTPError as e:
                hint = " — table missing; run scripts/create-state-tournament-results.sql first" if e.code == 404 else ""
                sys.exit(f"{method} {table} failed: {e.code} {e.read().decode()}{hint}")
            except urllib.error.URLError:
                time.sleep(2 * (attempt + 1))
        sys.exit(f"{method} {table} failed: network unavailable")

    def get_all(self, table, select, page=1000):
        out, offset = [], 0
        while True:
            rows = json.loads(self.request("GET", table, f"?{select}&limit={page}&offset={offset}"))
            out += rows
            offset += page
            if len(rows) < page:
                return out

    def division_filter(self, key):
        return "?" + "&".join(f"{k}=eq.{urllib.parse.quote(str(v))}" for k, v in zip(DIVISION_KEY, key))


def main():
    if "--confirm-only" in sys.argv:
        load_env()
        confirm_identities(Supabase())
        return
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 1:
        sys.exit(__doc__)
    path = os.path.expanduser(args[0])
    divisions, bouts, placers, errors = read_workbook(path)

    print(f"{os.path.basename(path)}: {len(divisions)} divisions, {len(bouts)} bouts, {len(placers)} placers")
    print(coverage(divisions, bouts))
    if errors:
        print(f"\n{len(errors)} problem(s) — nothing written:")
        print("\n".join(f"  {e}" for e in errors))
        sys.exit(1)
    if "--dry-run" in sys.argv:
        print("\nDry run — nothing written.")
        return

    load_env()
    db = Supabase()
    db.request("POST", "state_tournament_divisions", "?on_conflict=" + ",".join(DIVISION_KEY),
               list(divisions.values()), "resolution=merge-duplicates,return=minimal")
    for key in divisions:
        for table in ("state_tournament_placers", "state_tournament_bouts"):
            db.request("DELETE", table, db.division_filter(key), prefer="return=minimal")
    db.request("POST", "state_tournament_bouts", "", bouts, "return=minimal")
    db.request("POST", "state_tournament_placers", "", placers, "return=minimal")
    print("\nWritten.")
    if "--no-confirm" not in sys.argv:  # loading many states: confirm once at the end instead
        confirm_identities(db)


if __name__ == "__main__":
    main()
