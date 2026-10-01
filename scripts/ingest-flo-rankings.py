#!/usr/bin/env python3
"""
Find new FloWrestling HS rankings captures and import them into national_rankings.

Muse checks Flo's rankings every Monday and, when the tables changed, saves a dated workbook (the
"Flo 2026-27 HS Rankings.xlsx" layout: a P4P sheet, a By Weight sheet, a Source sheet carrying
"Last updated <date>"). This script, run weekly by a scheduled task, imports any capture it has
not seen before:

  1. looks in ~/Downloads, ~/Downloads/workspace/rankings/captures and ~/workspace/rankings/captures
     for .xlsx files whose name mentions Flo and rankings;
  2. skips files whose contents it already imported (by hash, so a renamed copy is not re-imported);
  3. writes the weight and P4P lists to scripts/data/flo-<YYYY-MM-DD>-{weight,p4p}.csv;
  4. runs scripts/import-national-rankings.ts for the month of Flo's "Last updated" date - a newer
     capture in the same month replaces that month's edition, which is what a weekly update means.

Prints one line per capture and a summary; exits 0 with "no new captures" when there is nothing.

  python3 scripts/ingest-flo-rankings.py [--dry-run]
"""

import csv
import datetime
import glob
import hashlib
import json
import os
import re
import subprocess
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HOME = os.path.expanduser("~")
FOLDERS = [
    os.path.join(HOME, "Downloads"),
    os.path.join(HOME, "Downloads", "workspace", "rankings", "captures"),
    os.path.join(HOME, "workspace", "rankings", "captures"),
]
STATE = os.path.join(HOME, ".recruitnc", "flo-rankings-imported.json")
# 2026-27 season: seniors graduate in 2027.
CLASS_YEAR = {"SR": 2027, "JR": 2028, "SO": 2029, "FR": 2030}
URL = "https://www.flowrestling.org/rankings"


def captures():
    found = []
    for folder in FOLDERS:
        for path in glob.glob(os.path.join(folder, "*.xlsx")):
            name = os.path.basename(path).lower()
            if "flo" in name and "rank" in name and not name.startswith("~$"):
                found.append(path)
    return sorted(set(found), key=os.path.getmtime)


def digest(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def last_updated(wb):
    """Flo's own 'Last updated <Mon D, YYYY>' from the Source sheet."""
    if "Source" not in wb.sheetnames:
        return None
    for row in wb["Source"].iter_rows(values_only=True):
        text = " ".join(str(c) for c in row if c)
        m = re.search(r"Last updated ([A-Z][a-z]+ \d{1,2}, \d{4})", text)
        if m:
            return datetime.datetime.strptime(m.group(1), "%b %d, %Y").date()
    return None


def rows_of(wb, sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = next((i for i, r in enumerate(rows) if r and "Rank" in [str(c) for c in r]), None)
    if header is None:
        return [], []
    return [str(c) for c in rows[header]], rows[header + 1:]


def lists(wb):
    weight, p4p = [], []
    cols, body = rows_of(wb, "By Weight")
    idx = {c: i for i, c in enumerate(cols)}
    for r in body:
        if r and r[idx["Rank"]] and r[idx["Name"]]:
            weight.append([r[idx["Rank"]], str(r[idx["Name"]]).strip(), r[idx["Weight"]],
                           CLASS_YEAR.get(str(r[idx["Year"]]).upper(), ""), str(r[idx["School"]] or "").strip(),
                           r[idx["State"]] or "", URL])
    cols, body = rows_of(wb, "P4P")
    idx = {c: i for i, c in enumerate(cols)}
    for r in body:
        if r and r[idx["Rank"]] and r[idx["Name"]]:
            p4p.append([r[idx["Rank"]], str(r[idx["Name"]]).strip(), r[idx["Weight"]],
                        CLASS_YEAR.get(str(r[idx["Year"]]).upper(), ""), str(r[idx["School"]] or "").strip(),
                        r[idx["State"]] or "", URL])
    return weight, p4p


def main():
    dry = "--dry-run" in sys.argv
    state = json.load(open(STATE)) if os.path.exists(STATE) else {}
    new = [p for p in captures() if digest(p) not in state]
    if not new:
        print("no new captures")
        return
    for path in new:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        updated = last_updated(wb)
        if not updated:
            print(f"SKIPPED {path}: no 'Last updated' date on a Source sheet")
            continue
        weight, p4p = lists(wb)
        if len(weight) < 100:
            print(f"SKIPPED {path}: only {len(weight)} weight rows - not the expected layout")
            continue
        print(f"{os.path.basename(path)}: Flo updated {updated}, {len(weight)} weight + {len(p4p)} P4P rows")
        if dry:
            continue
        for scope, data in (("weight", weight), ("p4p", p4p)):
            out = os.path.join(ROOT, "scripts", "data", f"flo-{updated}-{scope}.csv")
            with open(out, "w", newline="") as f:
                w = csv.writer(f)
                w.writerow(["rank", "athlete_name", "weight_class", "class_year", "high_school", "state", "source_url"])
                w.writerows(data)
            result = subprocess.run(
                ["npx", "tsx", "scripts/import-national-rankings.ts", "--file", out, "--source", "flowrestling",
                 "--month", updated.strftime("%Y-%m"), "--scope", scope],
                cwd=ROOT, capture_output=True, text=True,
            )
            if result.returncode != 0:
                sys.exit(f"import failed for {out}:\n{result.stdout}\n{result.stderr}")
            print("  " + " | ".join(l for l in result.stdout.splitlines() if l.startswith(("Wrote", "Granting", "Matched"))))
        state[digest(path)] = {"file": path, "flo_updated": str(updated), "imported_at": datetime.datetime.now().isoformat()}
        os.makedirs(os.path.dirname(STATE), exist_ok=True)
        json.dump(state, open(STATE, "w"), indent=2)
    print("done")


if __name__ == "__main__":
    main()
